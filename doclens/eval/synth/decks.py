"""Two synthetic PowerPoint decks (Japanese and English) with known facts, for the PPTX slice of the eval set (F5).

    python -m eval.synth.decks --out eval/data/decks

Each deck has a title, bullet text, a table and a chart. Ground truth is the list of key facts (strings that must be recovered from
the slides) and the table cells; the chart's values are listed too, for the chart → table experiment (F6).
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

DECKS = {
    "deck-ja": {
        "title": "2026年度 事業計画", "bullets": ["売上高は前年比12%増の1,234百万円", "営業利益は56百万円", "新製品は2026年10月24日に発売"],
        "table_title": "四半期別実績", "table": [["四半期", "売上高", "営業利益"], ["第1四半期", "300", "10"], ["第2四半期", "320", "14"], ["第3四半期", "310", "15"]],
        "chart": ("四半期別売上高", {"売上高": {"第1四半期": 300, "第2四半期": 320, "第3四半期": 310}}),
        "facts": ["2026年度 事業計画", "前年比12%増", "1,234百万円", "営業利益は56百万円", "2026年10月24日", "四半期別実績"],
    },
    "deck-en": {
        "title": "Product roadmap 2026", "bullets": ["Revenue grew 18% to $4.2M", "Launch of v2 planned for 14 March 2026", "Churn fell to 2.1%"],
        "table_title": "Quarterly results", "table": [["Quarter", "Revenue", "Profit"], ["Q1", "900", "80"], ["Q2", "1,050", "120"], ["Q3", "1,100", "130"]],
        "chart": ("Revenue by quarter", {"Revenue": {"Q1": 900, "Q2": 1050, "Q3": 1100}}),
        "facts": ["Product roadmap 2026", "grew 18% to $4.2M", "14 March 2026", "Churn fell to 2.1%", "Quarterly results"],
    },
}


def build(name: str, spec: dict, out: Path) -> Path:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = spec["title"]
    s1.placeholders[1].text = "\n".join(spec["bullets"])
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = spec["table_title"]
    rows, cols = len(spec["table"]), len(spec["table"][0])
    table = s2.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(8), Inches(0.5 * rows)).table
    for r, row in enumerate(spec["table"]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    title, series = spec["chart"]
    s3.shapes.title.text = title
    data = CategoryChartData()
    data.categories = list(next(iter(series.values())))
    for n, points in series.items():
        data.add_series(n, list(points.values()))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.8), Inches(8), Inches(4.5), data)
    path = out / f"{name}.pptx"
    prs.save(path)
    path.with_suffix(".json").write_text(json.dumps({"id": name, **{k: spec[k] for k in ("facts", "table")},
                                                      "chart": {"title": title, "series": series}}, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=Path(__file__).parent.parent / "data" / "decks")
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    print("wrote", ", ".join(build(n, s, a.out).name for n, s in DECKS.items()), "to", a.out)
